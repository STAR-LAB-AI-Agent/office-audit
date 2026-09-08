"""Read-only structural, field, and basic-format audits for .docx files.

The module intentionally keeps the deterministic audit path independent from
any model or network service.  It exposes :func:`audit_document` and
:func:`render_result` for tests and provides a small command-line interface
for course demonstrations.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.section import _Footer, _Header
from docx.shared import Length


SCHEMA_VERSION = "0.1"
SUPPORTED_MODES = ("full", "structure", "fields_format")
SUPPORTED_FORMATS = ("json", "markdown", "terminal")
SEVERITIES = ("error", "warning", "info")
DEFAULT_PLACEHOLDERS = (
    "TODO",
    "TBD",
    "待填写",
    "待补充",
    "请填写",
)
PLACEHOLDER_PATTERN = re.compile(r"TODO|TBD|待填写|待补充|请填写|_{3,}", re.IGNORECASE)
HEADING_PATTERN = re.compile(r"(?:heading|标题)\s*([1-9])", re.IGNORECASE)
UNSUPPORTED_OBJECT_TYPES = {
    "image_or_drawing": (
        "第一版只记录图片或绘图对象，不读取其中的文字",
        "图片中的标题、字段或扫描文字可能未参与审计",
        "请人工检查图片内容，或将需要审计的文字放入正文",
    ),
    "text_box": (
        "第一版不解析文本框或形状中的文字",
        "文本框中的标题、字段或占位内容可能未参与审计",
        "请人工检查文本框，或将文字移入正文",
    ),
    "shape": (
        "第一版只记录形状对象，不审计其内部语义",
        "形状中的文字或标注可能未参与审计",
        "请人工检查形状内容",
    ),
    "ole_object": (
        "第一版不读取嵌入式对象的内部内容",
        "嵌入的 Excel、PPT 或其他对象可能未参与审计",
        "请单独打开嵌入对象进行检查，并确认其许可证和来源",
    ),
    "smartart": (
        "第一版只识别 SmartArt 对象，不解析其文本和布局",
        "SmartArt 中的标题或字段可能未参与审计",
        "请人工检查 SmartArt，或将关键信息复制到正文",
    ),
    "comment": (
        "第一版不审计批注内容",
        "批注中的要求或待办项可能未参与审计",
        "请人工检查并处理批注",
    ),
    "tracked_change": (
        "第一版不合并或解释修订记录",
        "修订中的文字可能未按最终版本参与审计",
        "请在 Word 中确认修订状态后再提交",
    ),
}


class AuditInputError(Exception):
    """An expected input, rule, or output-path error."""

    def __init__(self, code: str, message: str, detail: str | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.detail = detail


_REQUEST_TERMS = {
    "full": ("全面", "完整", "综合", "全部", "全量", "full", "all"),
    "structure": ("标题", "层级", "章节", "目录", "结构", "heading", "section", "structure"),
    "fields_format": (
        "空字段",
        "空白项",
        "空单元格",
        "占位符",
        "格式",
        "字体",
        "字号",
        "对齐",
        "fields",
        "placeholder",
        "format",
    ),
}
_REQUEST_GENERAL_TERMS = ("审计", "检查", "文档", "word", "docx", "audit", "document")
_REQUEST_ONLY_TERMS = ("只", "仅", "only", "just")


def classify_request(request: str) -> dict[str, Any]:
    """Map a natural-language audit request to a deterministic CLI mode.

    This intentionally uses transparent keywords rather than a model.  A
    request that explicitly asks for two different narrow scopes is rejected
    as ambiguous instead of silently choosing one and hiding part of the
    user's intent.
    """

    text = str(request or "").strip().casefold()
    if not text:
        return {
            "status": "error",
            "error": {"code": "empty_request", "message": "自然语言请求不能为空"},
            "matched_terms": [],
        }

    matches = {
        group: [term for term in terms if term.casefold() in text]
        for group, terms in _REQUEST_TERMS.items()
    }
    has_full = bool(matches["full"])
    has_structure = bool(matches["structure"])
    has_fields = bool(matches["fields_format"])
    is_narrow = any(term.casefold() in text for term in _REQUEST_ONLY_TERMS)
    matched_terms = matches["full"] + matches["structure"] + matches["fields_format"]

    if has_structure and has_fields and is_narrow and not has_full:
        return {
            "status": "error",
            "error": {
                "code": "ambiguous_intent",
                "message": "请求同时限定了结构和字段/格式范围，请改用全面审计或明确指定一种范围",
            },
            "matched_terms": matched_terms,
        }
    if has_full or (has_structure and has_fields):
        return {
            "status": "ok",
            "intent": "full",
            "mode": "full",
            "matched_terms": matched_terms,
            "rationale": "请求表达了综合检查，映射到完整审计",
        }
    if has_structure:
        return {
            "status": "ok",
            "intent": "structure",
            "mode": "structure",
            "matched_terms": matched_terms,
            "rationale": "命中标题、层级或章节相关词，映射到结构审计",
        }
    if has_fields:
        return {
            "status": "ok",
            "intent": "fields_format",
            "mode": "fields_format",
            "matched_terms": matched_terms,
            "rationale": "命中空字段、占位符或格式相关词，映射到字段与格式审计",
        }
    if any(term.casefold() in text for term in _REQUEST_GENERAL_TERMS):
        return {
            "status": "ok",
            "intent": "full",
            "mode": "full",
            "matched_terms": [],
            "rationale": "未指定窄范围，按通用文档审计请求执行完整审计",
        }
    return {
        "status": "error",
        "error": {
            "code": "unknown_intent",
            "message": "无法从请求中识别审计范围，请说明结构、字段/格式或全面审计",
        },
        "matched_terms": matched_terms,
    }


class _Collector:
    def __init__(self, result: dict[str, Any]) -> None:
        self.result = result
        self._finding_number = 0
        self._unsupported_number = 0

    def finding(
        self,
        *,
        severity: str,
        rule_id: str,
        location: Mapping[str, Any],
        locator: Mapping[str, Any] | None = None,
        evidence: str,
        message: str,
        suggestion: str,
    ) -> None:
        if severity not in SEVERITIES:
            severity = "warning"
        self._finding_number += 1
        finding = {
            "id": f"F-{self._finding_number:03d}",
            "severity": severity,
            "rule_id": rule_id,
            "location": dict(location),
            "evidence": evidence,
            "message": message,
            "suggestion": suggestion,
        }
        if locator:
            finding["locator"] = dict(locator)
        self.result["findings"].append(finding)

    def unsupported(
        self,
        *,
        object_type: str,
        location: Mapping[str, Any],
        count: int = 1,
        reason: str | None = None,
        impact: str | None = None,
        suggestion: str | None = None,
    ) -> None:
        self._unsupported_number += 1
        defaults = UNSUPPORTED_OBJECT_TYPES.get(object_type, UNSUPPORTED_OBJECT_TYPES["shape"])
        self.result["unsupported_objects"].append(
            {
                "id": f"U-{self._unsupported_number:03d}",
                "severity": "warning",
                "object_type": object_type,
                "location": dict(location),
                "count": count,
                "reason": reason or defaults[0],
                "impact": impact or defaults[1],
                "suggestion": suggestion or defaults[2],
            }
        )


def _resolved_path(value: str | os.PathLike[str]) -> Path:
    return Path(value).expanduser().resolve(strict=False)


def _base_result(input_path: str | os.PathLike[str], mode: str) -> dict[str, Any]:
    path = _resolved_path(input_path)
    return {
        "schema_version": SCHEMA_VERSION,
        "audit_target": {
            "path": str(path),
            "format": path.suffix.lower().lstrip(".") or None,
        },
        "mode": mode,
        "scope": {
            "body": True,
            "tables": True,
            "headers_footers": True,
            "styles": True,
            "unsupported_detection": True,
        },
        "summary": {"error": 0, "warning": 0, "info": 0},
        "findings": [],
        "unsupported_objects": [],
        "metrics": {
            "elapsed_ms": 0,
            "input_chars_seen": 0,
            "paragraphs_seen": 0,
            "tables_seen": 0,
            "cells_seen": 0,
            "model_calls": 0,
        },
        "errors": [],
    }


def _add_error(result: dict[str, Any], code: str, message: str, detail: str | None = None) -> None:
    item: dict[str, Any] = {"code": code, "message": message}
    if detail:
        item["detail"] = detail
    result["errors"].append(item)


def _finalize(result: dict[str, Any], start_time: float) -> dict[str, Any]:
    counts = Counter(item.get("severity") for item in result["findings"])
    counts["error"] += len(result["errors"])
    result["summary"] = {severity: counts.get(severity, 0) for severity in SEVERITIES}
    result["metrics"]["elapsed_ms"] = max(0, round((time.perf_counter() - start_time) * 1000))
    return result


def _local_name(tag: Any) -> str:
    value = str(tag)
    return value.rsplit("}", 1)[-1]


def _xml_path(element: Any, root: Any) -> str:
    """Return a small structural path without exposing document text."""

    segments: list[str] = []
    current = element
    while current is not None:
        name = _local_name(getattr(current, "tag", "node"))
        parent = current.getparent() if hasattr(current, "getparent") else None
        if parent is None:
            segments.append(name)
            break
        same_name = [child for child in parent if _local_name(child.tag) == name]
        try:
            ordinal = same_name.index(current) + 1
        except ValueError:
            ordinal = 1
        segments.append(f"{name}[{ordinal}]")
        current = parent
        if current is root:
            segments.append(_local_name(getattr(root, "tag", "root")))
            break
    return "/".join(reversed(segments))


def _has_ancestor(element: Any, names: set[str], stop: Any) -> bool:
    current = element.getparent() if hasattr(element, "getparent") else None
    while current is not None and current is not stop:
        if _local_name(getattr(current, "tag", "")) in names:
            return True
        current = current.getparent() if hasattr(current, "getparent") else None
    return False


def _part_name(part: Any) -> str:
    owner = getattr(part, "part", part)
    return str(getattr(owner, "partname", "") or "")


def _iter_header_footer_parts(doc: Any) -> Iterable[tuple[str, int, str, Any]]:
    """Yield only header/footer parts actually referenced by the package.

    Accessing an absent ``Section.header`` variant through the high-level API
    can materialize a new part in memory.  Inspecting ``sectPr`` references
    avoids that side effect and prevents empty linked variants from becoming
    false audit targets.
    """

    seen: set[str] = set()
    variant_names = {
        "PRIMARY": "default",
        "FIRST_PAGE": "first_page",
        "EVEN_PAGE": "even_page",
    }
    for section_index, section in enumerate(doc.sections):
        sect_pr = getattr(section, "_sectPr", None)
        for part_kind in ("header", "footer"):
            references = getattr(sect_pr, f"{part_kind}Reference_lst", ())
            for reference in references:
                rel_id = getattr(reference, "rId", None)
                raw_part = section.part.related_parts.get(rel_id) if rel_id else None
                if raw_part is None:
                    continue
                partname = _part_name(raw_part)
                key = partname or f"{part_kind}:{section_index}:{rel_id}"
                if key in seen:
                    continue
                seen.add(key)
                wrapper_type = _Header if part_kind == "header" else _Footer
                part = wrapper_type(sect_pr, section.part, reference.type_)
                if getattr(part, "_element", None) is not None:
                    type_name = getattr(getattr(reference, "type_", None), "name", "")
                    variant = variant_names.get(type_name, type_name.lower() or "unknown")
                    yield part_kind, section_index, variant, part


def _paragraph_context(
    paragraph: Any,
    *,
    part: str,
    location: Mapping[str, Any],
) -> dict[str, Any]:
    return {
        "paragraph": paragraph,
        "text": paragraph.text or "",
        "part": part,
        "location": dict(location),
    }


def _collect_table(
    table: Any,
    *,
    part: str,
    table_index: list[int],
    paragraphs: list[dict[str, Any]],
    cells: list[dict[str, Any]],
    seen_cells: set[tuple[str, int]],
) -> None:
    current_table_index = table_index[0]
    table_index[0] += 1
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell_key = (part, id(getattr(cell, "_tc", cell)))
            if cell_key in seen_cells:
                continue
            seen_cells.add(cell_key)
            location = {
                "part": part,
                "table_index": current_table_index,
                "row_index": row_index,
                "cell_index": cell_index,
            }
            cells.append({"cell": cell, "text": cell.text or "", "location": location})
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                paragraph_location = dict(location)
                paragraph_location["paragraph_index"] = paragraph_index
                paragraphs.append(
                    _paragraph_context(paragraph, part=part, location=paragraph_location)
                )
            for nested in cell.tables:
                _collect_table(
                    nested,
                    part=part,
                    table_index=table_index,
                    paragraphs=paragraphs,
                    cells=cells,
                    seen_cells=seen_cells,
                )


def _collect_content(doc: Any, result: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    paragraphs: list[dict[str, Any]] = []
    cells: list[dict[str, Any]] = []
    seen_cells: set[tuple[str, int]] = set()
    table_index = [0]

    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        paragraphs.append(
            _paragraph_context(
                paragraph,
                part="body",
                location={"part": "body", "paragraph_index": paragraph_index},
            )
        )
    for table in doc.tables:
        _collect_table(
            table,
            part="table",
            table_index=table_index,
            paragraphs=paragraphs,
            cells=cells,
            seen_cells=seen_cells,
        )

    for part_kind, section_index, variant, part in _iter_header_footer_parts(doc):
        for paragraph_index, paragraph in enumerate(part.paragraphs):
            paragraphs.append(
                _paragraph_context(
                    paragraph,
                    part=part_kind,
                    location={
                        "part": part_kind,
                        "section_index": section_index,
                        "variant": variant,
                        "paragraph_index": paragraph_index,
                    },
                )
            )
        for table in part.tables:
            _collect_table(
                table,
                part=part_kind,
                table_index=table_index,
                paragraphs=paragraphs,
                cells=cells,
                seen_cells=seen_cells,
            )

    result["metrics"]["paragraphs_seen"] = len(paragraphs)
    result["metrics"]["tables_seen"] = table_index[0]
    result["metrics"]["cells_seen"] = len(cells)
    result["metrics"]["input_chars_seen"] = sum(len(item["text"]) for item in paragraphs)
    for index, context in enumerate(paragraphs):
        context["locator"] = _paragraph_locator(paragraphs, index)
    return paragraphs, cells


def _style_name(paragraph: Any) -> str:
    try:
        return paragraph.style.name or ""
    except (AttributeError, ValueError):
        return ""


def _heading_level(paragraph: Any) -> int | None:
    style_name = _style_name(paragraph)
    match = HEADING_PATTERN.search(style_name)
    return int(match.group(1)) if match else None


def _compact_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def _preview(value: str, limit: int = 80) -> str:
    text = _compact_text(value)
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _location_label(location: Mapping[str, Any]) -> str:
    part = str(location.get("part", "文档"))
    paragraph_index = location.get("paragraph_index")
    table_index = location.get("table_index")
    if table_index is not None:
        story = {"table": "正文", "header": "页眉", "footer": "页脚"}.get(part, part)
        label = (
            f"{story}第{int(table_index) + 1}个表格"
            f"第{int(location.get('row_index', 0)) + 1}行"
            f"第{int(location.get('cell_index', 0)) + 1}列"
        )
        if paragraph_index is not None:
            label += f"内第{int(paragraph_index) + 1}段"
        return label
    if part == "body":
        if paragraph_index is not None:
            return f"正文第{int(paragraph_index) + 1}段"
        return "正文"
    if part in {"header", "footer"}:
        story = "页眉" if part == "header" else "页脚"
        variant = {
            "default": "默认",
            "first_page": "首页",
            "even_page": "偶数页",
        }.get(str(location.get("variant", "")), "")
        section = int(location.get("section_index", 0)) + 1
        label = f"第{section}节{variant}{story}"
        if paragraph_index is not None:
            label += f"第{int(paragraph_index) + 1}段"
        return label
    return part


def _paragraph_container_key(location: Mapping[str, Any]) -> tuple[tuple[str, Any], ...]:
    return tuple(sorted((key, value) for key, value in location.items() if key != "paragraph_index"))


def _paragraph_locator(
    paragraphs: Sequence[Mapping[str, Any]], index: int
) -> dict[str, Any]:
    context = paragraphs[index]
    location = context["location"]
    container = _paragraph_container_key(location)
    locator: dict[str, Any] = {"label": _location_label(location)}
    text = _preview(str(context["text"]))
    if text:
        locator["text"] = text

    for nearby_index in range(index - 1, -1, -1):
        nearby = paragraphs[nearby_index]
        if _paragraph_container_key(nearby["location"]) != container:
            continue
        nearby_text = _preview(str(nearby["text"]))
        if nearby_text:
            locator["previous_text"] = nearby_text
            break
    for nearby_index in range(index + 1, len(paragraphs)):
        nearby = paragraphs[nearby_index]
        if _paragraph_container_key(nearby["location"]) != container:
            continue
        nearby_text = _preview(str(nearby["text"]))
        if nearby_text:
            locator["next_text"] = nearby_text
            break
    return locator


def _length_value(value: Length | None) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value.pt), 2)
    except (AttributeError, TypeError, ValueError):
        return None


def _format_signature(paragraph: Any) -> tuple[Any, ...]:
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    font = getattr(run, "font", None)
    fmt = paragraph.paragraph_format
    return (
        _style_name(paragraph),
        getattr(font, "name", None),
        _length_value(getattr(font, "size", None)),
        getattr(font, "bold", None),
        getattr(font, "italic", None),
        str(paragraph.alignment) if paragraph.alignment is not None else None,
        _length_value(getattr(fmt, "left_indent", None)),
        _length_value(getattr(fmt, "right_indent", None)),
        _length_value(getattr(fmt, "first_line_indent", None)),
        _length_value(getattr(fmt, "space_before", None)),
        _length_value(getattr(fmt, "space_after", None)),
        fmt.line_spacing,
    )


_FORMAT_FIELD_NAMES = (
    "样式",
    "字体",
    "字号",
    "加粗",
    "斜体",
    "对齐",
    "左缩进",
    "右缩进",
    "首行缩进",
    "段前间距",
    "段后间距",
    "行距",
)


def _format_value(value: Any) -> str:
    if value is None:
        return "继承样式/未显式设置"
    if value is True:
        return "是"
    if value is False:
        return "否"
    return str(value)


def _format_differences(current: Sequence[Any], dominant: Sequence[Any]) -> str:
    differences = []
    for name, current_value, dominant_value in zip(_FORMAT_FIELD_NAMES, current, dominant):
        if current_value != dominant_value:
            differences.append(
                f"{name}：主导={_format_value(dominant_value)}，本段={_format_value(current_value)}"
            )
    return "；".join(differences) or "未识别到可展示的差异"


def _rule_severity(rules: Mapping[str, Any], rule_id: str, default: str) -> str:
    overrides = rules.get("severity_overrides", {})
    value = overrides.get(rule_id) if isinstance(overrides, Mapping) else None
    return value if value in SEVERITIES else default


def _placeholder_regex(rules: Mapping[str, Any]) -> re.Pattern[str]:
    configured = rules.get("placeholders")
    if configured is None:
        return PLACEHOLDER_PATTERN
    if isinstance(configured, str):
        configured = [configured]
    if not isinstance(configured, Sequence) or isinstance(configured, (bytes, bytearray)):
        return PLACEHOLDER_PATTERN
    values = [str(item) for item in configured if str(item)]
    if not values:
        return re.compile(r"(?!x)x")
    escaped = "|".join(re.escape(item) for item in values)
    return re.compile(escaped, re.IGNORECASE)


def _audit_structure(
    paragraphs: Sequence[Mapping[str, Any]],
    collector: _Collector,
    *,
    required_sections: Sequence[str] | None,
    rules: Mapping[str, Any],
) -> None:
    headings: list[dict[str, Any]] = []
    for context in paragraphs:
        level = _heading_level(context["paragraph"])
        if level is not None and _compact_text(context["text"]):
            item = dict(context)
            item["level"] = level
            headings.append(item)

    if not headings:
        collector.finding(
            severity=_rule_severity(rules, "structure.no_headings", "info"),
            rule_id="structure.no_headings",
            location={"part": "body"},
            evidence="未识别到 Heading/标题样式段落",
            message="文档中未识别到可用于结构审计的标题样式",
            suggestion="如果文档确实有章节，请使用 Word 标题样式；否则提供自定义结构规则",
        )

    previous_level: int | None = None
    previous_title = ""
    for context in headings:
        level = context["level"]
        if previous_level is not None and level > previous_level + 1:
            collector.finding(
                severity=_rule_severity(rules, "structure.heading_level_jump", "warning"),
                rule_id="structure.heading_level_jump",
                location=context["location"],
                locator=context.get("locator"),
                evidence=f"前一标题层级为 {previous_level}，当前层级为 {level}",
                message="标题层级存在跳级",
                suggestion="检查是否缺少中间层级，或在规则配置中声明允许的层级关系",
            )
        previous_level = level
        previous_title = _preview(context["text"])

    if required_sections:
        heading_texts = [_normalized_text(item["text"]) for item in headings]
        for required in required_sections:
            wanted = _normalized_text(str(required))
            if not wanted:
                continue
            if not any(wanted in title for title in heading_texts):
                collector.finding(
                    severity=_rule_severity(rules, "structure.required_section_missing", "error"),
                    rule_id="structure.required_section_missing",
                    location={"part": "body", "section_requirement": str(required)},
                    evidence=f"未在标题样式中找到“{_preview(str(required))}”",
                    message="用户指定的必需章节缺失",
                    suggestion="补充该章节，或确认章节名称和规则清单是否正确",
                )
    else:
        collector.finding(
            severity=_rule_severity(rules, "structure.required_sections_not_specified", "info"),
            rule_id="structure.required_sections_not_specified",
            location={"part": "body"},
            evidence="调用时没有提供 required_sections",
            message="未指定标准：未提供必需章节清单，未按固定模板判断缺失章节",
            suggestion="如有课程、单位或用户规定，请通过参数或规则文件提供必需章节清单",
        )


def _audit_fields_and_format(
    paragraphs: Sequence[Mapping[str, Any]],
    cells: Sequence[Mapping[str, Any]],
    collector: _Collector,
    *,
    rules: Mapping[str, Any],
) -> None:
    placeholder_pattern = _placeholder_regex(rules)
    for context in paragraphs:
        text = str(context["text"])
        location = context["location"]
        is_top_level_body = context["part"] == "body" and "table_index" not in location
        if not _compact_text(text) and is_top_level_body:
            collector.finding(
                severity=_rule_severity(rules, "fields.empty_paragraph", "warning"),
                rule_id="fields.empty_paragraph",
                location=location,
                locator=context.get("locator"),
                evidence="该正文段落不包含可见文字；请结合前后文定位",
                message="发现正文空白段落标记",
                suggestion="在 Word 中开启“显示/隐藏编辑标记(¶)”核对；若只是有意留白可忽略，否则删除多余段落",
            )
        match = placeholder_pattern.search(text) if "table_index" not in location else None
        if match:
            collector.finding(
                severity=_rule_severity(rules, "fields.placeholder", "warning"),
                rule_id="fields.placeholder",
                location=location,
                locator=context.get("locator"),
                evidence=f"命中占位符“{_preview(match.group(0))}”",
                message="发现可能未完成的占位内容",
                suggestion="提交前替换占位符，或在规则配置中明确说明该占位符允许保留",
            )

    for cell in cells:
        text = str(cell["text"])
        if not _compact_text(text):
            collector.finding(
                severity=_rule_severity(rules, "fields.empty_cell", "warning"),
                rule_id="fields.empty_cell",
                location=cell["location"],
                evidence="单元格不包含可见文字",
                message="发现空表格单元格",
                suggestion="确认该单元格是否应填写，或在规则中标记为允许为空",
            )
        match = placeholder_pattern.search(text)
        if match:
            collector.finding(
                severity=_rule_severity(rules, "fields.placeholder_cell", "warning"),
                rule_id="fields.placeholder_cell",
                location=cell["location"],
                evidence=f"单元格命中占位符“{_preview(match.group(0))}”",
                message="发现表格中的可能未完成占位内容",
                suggestion="提交前替换占位符，或在规则配置中明确说明该占位符允许保留",
            )

    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for context in paragraphs:
        if not _compact_text(str(context["text"])):
            continue
        if _heading_level(context["paragraph"]) is not None:
            continue
        groups[(str(context["part"]), _style_name(context["paragraph"]))].append(dict(context))

    for group in groups.values():
        if len(group) < 3:
            continue
        signatures = Counter(_format_signature(item["paragraph"]) for item in group)
        if len(signatures) < 2:
            continue
        dominant, dominant_count = signatures.most_common(1)[0]
        if dominant_count < 2:
            continue
        for context in group:
            current = _format_signature(context["paragraph"])
            if current != dominant:
                differences = _format_differences(current, dominant)
                collector.finding(
                    severity=_rule_severity(rules, "format.dominant_style_outlier", "warning"),
                    rule_id="format.dominant_style_outlier",
                    location=context["location"],
                    locator=context.get("locator"),
                    evidence=f"该段落的基础格式偏离同类段落主导格式。格式差异：{differences}",
                    message="发现同类段落的基础格式离群",
                    suggestion="对照同类段落检查字体、字号、对齐、缩进和间距；如有意不同，请通过规则说明",
                )


def _unsupported_location(part: str, root: Any, element: Any, ordinal: int) -> dict[str, Any]:
    return {
        "part": part,
        "xml_path": _xml_path(element, root),
        "object_index": ordinal,
    }


def _scan_xml_part(root: Any, part: str, collector: _Collector) -> int:
    if root is None:
        return 0
    counts: Counter[str] = Counter()
    ordinals: Counter[str] = Counter()
    for element in root.iter():
        local = _local_name(getattr(element, "tag", ""))
        object_type: str | None = None
        if local == "drawing":
            child_names = {_local_name(getattr(child, "tag", "")) for child in element.iter()}
            if {"relIds", "diagram"} & child_names:
                object_type = "smartart"
            elif "txbxContent" in child_names or "txbx" in child_names:
                object_type = "text_box"
            elif "wsp" in child_names:
                object_type = "shape"
            else:
                object_type = "image_or_drawing"
        elif local == "pict":
            child_names = {_local_name(getattr(child, "tag", "")) for child in element.iter()}
            object_type = "text_box" if "textbox" in child_names else "image_or_drawing"
        elif local == "txbxContent" and not _has_ancestor(element, {"drawing", "pict"}, root):
            object_type = "text_box"
        elif local in {"object", "OLEObject"}:
            object_type = "ole_object"
        elif local == "commentReference":
            object_type = "comment"
        elif local in {
            "ins",
            "del",
            "moveFrom",
            "moveTo",
            "moveFromRangeStart",
            "moveFromRangeEnd",
            "moveToRangeStart",
            "moveToRangeEnd",
        }:
            object_type = "tracked_change"
        if object_type is None:
            continue
        counts[object_type] += 1
        ordinals[object_type] += 1
        collector.unsupported(
            object_type=object_type,
            location=_unsupported_location(part, root, element, ordinals[object_type]),
            count=1,
        )
    return counts.get("comment", 0)


def _scan_package_comments(input_path: Path, collector: _Collector, existing_references: int) -> None:
    try:
        with zipfile.ZipFile(input_path) as archive:
            names = [name for name in archive.namelist() if name.lower().endswith("comments.xml")]
            for name in names:
                payload = archive.read(name)
                comment_count = len(re.findall(rb"<w:comment(?:\s|>)", payload))
                if comment_count <= existing_references:
                    continue
                for index in range(existing_references, comment_count):
                    collector.unsupported(
                        object_type="comment",
                        location={"part": "comments", "package_part": name, "comment_index": index},
                        count=1,
                    )
    except (OSError, zipfile.BadZipFile, KeyError):
        # Opening the main document already reports the stable input error;
        # package-level comment inspection is only an enhancement.
        return


def _scan_unsupported(doc: Any, input_path: Path, collector: _Collector) -> None:
    comment_references = 0
    body = getattr(getattr(doc, "element", None), "body", None)
    comment_references += _scan_xml_part(body, "body", collector)
    seen_parts: set[str] = set()
    for part_kind, section_index, variant, part in _iter_header_footer_parts(doc):
        partname = _part_name(part)
        key = partname or f"{part_kind}:{section_index}:{variant}"
        if key in seen_parts:
            continue
        seen_parts.add(key)
        comment_references += _scan_xml_part(getattr(part, "_element", None), part_kind, collector)
    _scan_package_comments(input_path, collector, comment_references)


def _normalize_rules(rules: Mapping[str, Any] | None) -> dict[str, Any]:
    if rules is None:
        return {}
    if not isinstance(rules, Mapping):
        raise AuditInputError("invalid_rules", "规则配置必须是 JSON 对象")
    return dict(rules)


def audit_document(
    input_path: str | os.PathLike[str],
    *,
    mode: str = "full",
    required_sections: Sequence[str] | None = None,
    rules: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Audit a .docx without modifying it and return the stable result model."""

    start_time = time.perf_counter()
    result = _base_result(input_path, mode)
    collector = _Collector(result)
    path = _resolved_path(input_path)
    try:
        if mode not in SUPPORTED_MODES:
            raise AuditInputError("invalid_mode", f"不支持的审计模式：{mode}")
        if not path.exists():
            raise AuditInputError("input_not_found", "输入文件不存在", str(path))
        if not path.is_file():
            raise AuditInputError("input_not_file", "输入路径不是文件", str(path))
        if path.suffix.lower() != ".docx":
            raise AuditInputError("unsupported_format", "第一版只接受 .docx 文件", path.suffix or "无扩展名")
        if not os.access(path, os.R_OK):
            raise AuditInputError("input_not_readable", "输入文件不可读取", str(path))
        normalized_rules = _normalize_rules(rules)
        if required_sections is None:
            configured = normalized_rules.get("required_sections")
            if isinstance(configured, Sequence) and not isinstance(configured, (str, bytes, bytearray)):
                required_sections = [str(item) for item in configured]
        try:
            doc = Document(str(path))
        except (PackageNotFoundError, zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
            raise AuditInputError("invalid_docx", "文件不是可解析的有效 .docx 文档", type(exc).__name__) from None

        paragraphs, cells = _collect_content(doc, result)
        if mode in {"full", "structure"}:
            _audit_structure(
                paragraphs,
                collector,
                required_sections=required_sections,
                rules=normalized_rules,
            )
        if mode in {"full", "fields_format"}:
            _audit_fields_and_format(paragraphs, cells, collector, rules=normalized_rules)
        _scan_unsupported(doc, path, collector)
    except AuditInputError as exc:
        _add_error(result, exc.code, exc.message, exc.detail)
    except Exception as exc:  # pragma: no cover - defensive CLI boundary
        _add_error(result, "internal_error", "审计过程发生内部错误", type(exc).__name__)
    return _finalize(result, start_time)


def _location_text(location: Mapping[str, Any]) -> str:
    return json.dumps(dict(location), ensure_ascii=False, separators=(",", ":"))


def render_result(result: Mapping[str, Any], format_name: str = "json") -> str:
    """Render one result model as JSON or a human-readable Markdown projection."""

    if format_name not in SUPPORTED_FORMATS:
        raise ValueError(f"unsupported output format: {format_name}")
    if format_name == "json":
        return json.dumps(result, ensure_ascii=False, indent=2) + "\n"

    target = result.get("audit_target", {})
    summary = result.get("summary", {})
    lines = [
        "# AI Office 文档审计",
        f"- 文件：{target.get('path', '')}",
        f"- 模式：{result.get('mode', '')}",
        f"- 统计：error={summary.get('error', 0)}，warning={summary.get('warning', 0)}，info={summary.get('info', 0)}",
        "",
    ]
    routing = result.get("request_routing")
    if isinstance(routing, Mapping):
        lines.insert(
            3,
            f"- 意图：{routing.get('intent', '未识别')}（匹配：{', '.join(routing.get('matched_terms', [])) or '通用审计'}）",
        )
    errors = result.get("errors", [])
    if errors:
        lines.append("## 输入或运行错误")
        for error in errors:
            detail = f"（{error['detail']}）" if error.get("detail") else ""
            lines.append(f"- [{error.get('code', 'error')}] {error.get('message', '')}{detail}")
        lines.append("")

    lines.append("## 审计问题")
    findings = result.get("findings", [])
    if not findings:
        lines.append("- 未发现规则问题。")
    for finding in findings:
        location = finding.get("location", {})
        locator = finding.get("locator", {})
        human_location = (
            locator.get("label") if isinstance(locator, Mapping) else None
        ) or _location_label(location)
        lines.extend(
            [
                f"### {finding.get('id', '')} [{finding.get('severity', '')}] {finding.get('message', '')}",
                f"- 规则：{finding.get('rule_id', '')}",
                f"- 位置：{human_location}",
            ]
        )
        if isinstance(locator, Mapping) and locator.get("text"):
            lines.append(f"- 段落文字：“{locator['text']}”")
        nearby = []
        if isinstance(locator, Mapping) and locator.get("previous_text"):
            nearby.append(f"前文：“{locator['previous_text']}”")
        if isinstance(locator, Mapping) and locator.get("next_text"):
            nearby.append(f"后文：“{locator['next_text']}”")
        if nearby:
            lines.append(f"- 附近文字：{'；'.join(nearby)}")
        lines.extend(
            [
                f"- 技术定位：{_location_text(location)}",
                f"- 证据：{finding.get('evidence', '')}",
                f"- 建议：{finding.get('suggestion', '')}",
                "",
            ]
        )

    lines.append("## 未审计对象")
    unsupported = result.get("unsupported_objects", [])
    if not unsupported:
        lines.append("- 未识别到基础范围之外的对象。")
    for item in unsupported:
        location = item.get("location", {})
        lines.extend(
            [
                f"### {item.get('id', '')} [{item.get('object_type', '')}]",
                f"- 位置：{_location_label(location)}",
                f"- 技术定位：{_location_text(location)}",
                f"- 数量：{item.get('count', 0)}",
                f"- 原因：{item.get('reason', '')}",
                f"- 影响：{item.get('impact', '')}",
                f"- 建议：{item.get('suggestion', '')}",
                "",
            ]
        )
    return "\n".join(lines).rstrip() + "\n"


def _load_json_mapping(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        raise AuditInputError("rules_not_found", "规则文件不存在", str(path)) from None
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AuditInputError("invalid_rules", "规则文件不是可解析的 JSON", type(exc).__name__) from None


def _load_rules(path_value: str | None) -> dict[str, Any]:
    if not path_value:
        return {}
    path = _resolved_path(path_value)
    value = _load_json_mapping(path)
    if not isinstance(value, Mapping):
        raise AuditInputError("invalid_rules", "规则文件顶层必须是 JSON 对象")
    return dict(value)


def _load_required_sections(path_value: str | None) -> list[str] | None:
    if not path_value:
        return None
    path = _resolved_path(path_value)
    if not path.exists():
        raise AuditInputError("sections_not_found", "必需章节文件不存在", str(path))
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        raise AuditInputError("invalid_sections", "必需章节文件不可读取", type(exc).__name__) from None
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        return [line.strip() for line in text.splitlines() if line.strip()]
    if isinstance(value, Mapping):
        value = value.get("required_sections", value.get("sections"))
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes, bytearray)):
        raise AuditInputError("invalid_sections", "必需章节文件应为 JSON 数组或逐行文本")
    return [str(item).strip() for item in value if str(item).strip()]


def _same_path(left: Path, right: Path) -> bool:
    return os.path.normcase(str(left.resolve(strict=False))) == os.path.normcase(
        str(right.resolve(strict=False))
    )


def _emit(result: Mapping[str, Any], format_name: str, output_value: str | None) -> None:
    rendered = render_result(result, format_name)
    if output_value:
        output_path = _resolved_path(output_value)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(rendered, encoding="utf-8")
        return
    sys.stdout.write(rendered)


def _append_audit_log(result: Mapping[str, Any], log_path: Path) -> None:
    """Append one metadata-only JSONL record without document evidence."""

    errors = result.get("errors", [])
    record = {
        "event": "audit_completed",
        "schema_version": result.get("schema_version"),
        "mode": result.get("mode"),
        "summary": dict(result.get("summary", {})),
        "metrics": dict(result.get("metrics", {})),
        "error_codes": [item.get("code") for item in errors if isinstance(item, Mapping)],
    }
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Audit a .docx document without modifying the source file.")
    parser.add_argument("--input", required=True, help="path to the .docx document")
    parser.add_argument("--request", help="natural-language request used to infer the audit mode")
    parser.add_argument("--mode", choices=SUPPORTED_MODES, default=None)
    parser.add_argument("--required-sections", help="JSON/text file containing required section names")
    parser.add_argument("--rules", help="JSON rule configuration")
    parser.add_argument("--format", choices=SUPPORTED_FORMATS, default="json", dest="format_name")
    parser.add_argument("--output", help="separate report path; never overwrite --input")
    parser.add_argument("--log", help="optional metadata-only JSONL log path")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    result: dict[str, Any]
    selected_mode = args.mode or "full"
    routing: dict[str, Any] | None = None
    try:
        if args.request is not None:
            routing = classify_request(args.request)
            if routing.get("status") != "ok":
                result = _base_result(args.input, selected_mode)
                result["request_routing"] = routing
                error = routing.get("error", {})
                _add_error(
                    result,
                    str(error.get("code", "invalid_request")),
                    str(error.get("message", "无法识别自然语言请求")),
                )
                result = _finalize(result, time.perf_counter())
                _emit(result, args.format_name, None)
                return 2
            selected_mode = args.mode or str(routing["mode"])
            routing = dict(routing)
            routing["selected_mode"] = selected_mode
        rules = _load_rules(args.rules)
        sections = _load_required_sections(args.required_sections)
        input_path = _resolved_path(args.input)
        output_path = _resolved_path(args.output) if args.output else None
        log_path = _resolved_path(args.log) if args.log else None
        if (
            (output_path is not None and _same_path(input_path, output_path))
            or (log_path is not None and _same_path(input_path, log_path))
            or (output_path is not None and log_path is not None and _same_path(output_path, log_path))
        ):
            result = _base_result(input_path, selected_mode)
            if routing is not None:
                result["request_routing"] = routing
            _add_error(result, "unsafe_output_path", "报告或日志路径不能覆盖输入文档或彼此覆盖")
            result = _finalize(result, time.perf_counter())
            _emit(result, args.format_name, None)
            return 2
        result = audit_document(
            input_path,
            mode=selected_mode,
            required_sections=sections,
            rules=rules,
        )
        if routing is not None:
            result["request_routing"] = routing
        if log_path is not None:
            try:
                _append_audit_log(result, log_path)
            except OSError as exc:
                _add_error(result, "log_output_error", "日志写入失败", type(exc).__name__)
                result["summary"]["error"] += 1
        _emit(result, args.format_name, str(output_path) if output_path else None)
        return 2 if result.get("errors") else 0
    except AuditInputError as exc:
        result = _base_result(args.input, selected_mode)
        if routing is not None:
            result["request_routing"] = routing
        _add_error(result, exc.code, exc.message, exc.detail)
        result = _finalize(result, time.perf_counter())
        _emit(result, args.format_name, None)
        return 2
    except OSError as exc:
        result = _base_result(args.input, selected_mode)
        if routing is not None:
            result["request_routing"] = routing
        _add_error(result, "output_error", "报告输出失败", type(exc).__name__)
        result = _finalize(result, time.perf_counter())
        _emit(result, args.format_name, None)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "audit_document",
    "build_parser",
    "classify_request",
    "main",
    "render_result",
]
