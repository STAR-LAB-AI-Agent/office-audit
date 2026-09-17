"""Generate temporary controlled .docx cases and verify expected audit labels."""

from __future__ import annotations

import argparse
import json
import re
import tempfile
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

try:
    from .audit_docx import audit_document
except ImportError:  # direct execution from the repository root
    from audit_docx import audit_document


Builder = Callable[[Document], None]


def _clean_general_document(document: Document) -> None:
    document.add_heading("项目说明", level=1)
    document.add_paragraph("这是一份用于规则基线的通用项目说明。")
    document.add_paragraph("正文内容完整，便于验证正常文档不会产生警告。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "状态"
    table.cell(1, 1).text = "已完成"


def _heading_level_jump(document: Document) -> None:
    document.add_heading("项目说明", level=1)
    document.add_heading("实现细节直接使用三级标题", level=3)
    document.add_paragraph("这里仅用于验证结构规则。")


def _required_section_missing(document: Document) -> None:
    document.add_heading("项目说明", level=1)
    document.add_paragraph("用户要求的风险章节没有出现在该文档中。")


def _empty_placeholder_and_format(document: Document) -> None:
    document.add_heading("字段与格式", level=1)
    document.add_paragraph("同类正文一")
    document.add_paragraph("同类正文二")
    outlier = document.add_paragraph("同类正文格式离群")
    outlier.runs[0].bold = True
    document.add_paragraph("")
    document.add_paragraph("TODO：提交前补充本字段。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = ""


def _unsupported_drawing_object(document: Document) -> None:
    document.add_heading("对象检查", level=1)
    document.add_paragraph("正文仍可读取，但绘图对象内部文字不纳入第一版语义审计。")
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run._r.append(OxmlElement("w:drawing"))


def _clean_rich_layout(document: Document) -> None:
    document.add_heading("第一章 总体概述", level=1)
    document.add_paragraph("本章节介绍项目的基本背景与设计目标。")
    document.add_heading("1.1 核心需求", level=2)
    document.add_paragraph("需求列表如下：")
    document.add_paragraph("支持多宿主调用", style="List Bullet")
    document.add_paragraph("支持只读审计", style="List Bullet")
    document.add_paragraph("这是一段关键引用规范说明。", style="Quote")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "参数名"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "input_path"
    table.cell(1, 1).text = "输入文件绝对路径"
    table.cell(2, 0).text = "mode"
    table.cell(2, 1).text = "审计模式"


def _single_empty_paragraph(document: Document) -> None:
    document.add_heading("单空段测试", level=1)
    document.add_paragraph("上文正文段落")
    document.add_paragraph("")
    document.add_paragraph("下文正文段落")


def _consecutive_empty_paragraphs(document: Document) -> None:
    document.add_heading("连续空段测试", level=1)
    document.add_paragraph("上文正文段落")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("下文正文段落")


def _table_boundary_blanks(document: Document) -> None:
    document.add_heading("跨表格边界留白", level=1)
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表头1"
    table.cell(0, 1).text = "表头2"
    table.cell(1, 0).text = "值1"
    table.cell(1, 1).text = "值2"
    document.add_paragraph("")
    document.add_paragraph("表格后正文")


def _content_carrier_paragraphs(document: Document) -> None:
    document.add_heading("内容承载段落", level=1)
    document.add_paragraph("对象上文")
    p1 = document.add_paragraph()
    p1.add_run()._r.append(OxmlElement("w:drawing"))
    p2 = document.add_paragraph()
    p2._p.append(OxmlElement("m:oMathPara"))
    document.add_page_break()
    document.add_paragraph("对象下文")


def _visual_title_candidate(document: Document) -> None:
    document.add_paragraph("项目说明综合方案", style="Title")
    document.add_paragraph("这是一篇测试视觉标题候选与必需章节判定的文档。")


def _table_column_baselines(document: Document) -> None:
    table = document.add_table(rows=5, cols=2)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = "内容"
            if j == 1 or i == 0:
                cell.paragraphs[0].runs[0].bold = True


def _table_cell_outlier(document: Document) -> None:
    table = document.add_table(rows=5, cols=2)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = "内容"
            if j == 1 or i == 0:
                cell.paragraphs[0].runs[0].bold = True
    table.cell(3, 0).paragraphs[0].runs[0].italic = True


def _no_majority_format(document: Document) -> None:
    document.add_heading("无多数格式样例", level=1)
    p0 = document.add_paragraph("第一种格式段落")
    p0.runs[0].bold = True
    p1 = document.add_paragraph("第二种格式段落")
    p1.runs[0].italic = True
    p2 = document.add_paragraph("第三种格式段落")
    p2.runs[0].font.size = Pt(14)
    p3 = document.add_paragraph("第四种格式段落")
    p3.runs[0].font.name = "Courier New"


def _unsupported_mixed_objects(document: Document) -> None:
    document.add_heading("复杂对象样例", level=1)
    p1 = document.add_paragraph("正文文字")
    p1.add_run()._r.append(OxmlElement("w:commentReference"))
    p2 = document.add_paragraph()
    p2._p.append(OxmlElement("w:txbxContent"))
    p3 = document.add_paragraph()
    p3.add_run()._r.append(OxmlElement("w:object"))


BUILDERS: dict[str, Builder] = {
    "clean_general_document": _clean_general_document,
    "heading_level_jump": _heading_level_jump,
    "required_section_missing": _required_section_missing,
    "empty_placeholder_and_format": _empty_placeholder_and_format,
    "unsupported_drawing_object": _unsupported_drawing_object,
    "clean_rich_layout": _clean_rich_layout,
    "single_empty_paragraph": _single_empty_paragraph,
    "consecutive_empty_paragraphs": _consecutive_empty_paragraphs,
    "table_boundary_blanks": _table_boundary_blanks,
    "content_carrier_paragraphs": _content_carrier_paragraphs,
    "visual_title_candidate": _visual_title_candidate,
    "table_column_baselines": _table_column_baselines,
    "table_cell_outlier": _table_cell_outlier,
    "no_majority_format": _no_majority_format,
    "unsupported_mixed_objects": _unsupported_mixed_objects,
}

DEFECT_SEVERITIES = {"warning", "error"}


def _load_manifest(path_or_mapping: Path | Mapping[str, Any]) -> Mapping[str, Any]:
    if isinstance(path_or_mapping, Mapping):
        value = path_or_mapping
    else:
        value = json.loads(Path(path_or_mapping).read_text(encoding="utf-8"))
    if not isinstance(value, Mapping) or not isinstance(value.get("cases"), list) or not value["cases"]:
        raise ValueError("受控样例清单必须是包含 cases 数组的 JSON 对象")
    ids = set()
    for case in value["cases"]:
        if not isinstance(case, Mapping):
            raise ValueError("每个样例必须是 JSON 对象")
        case_id = case.get("id")
        if not isinstance(case_id, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", case_id):
            raise ValueError("样例 id 只能包含字母数字下划线或连字符，且不超过80字符")
        if case_id in ids:
            raise ValueError("样例 id 不得重复")
        ids.add(case_id)
        expected = case.get("expected_findings", [])
        absent = case.get("expected_absent_rules", [])
        if not isinstance(expected, list) or any(not isinstance(e, Mapping) or not isinstance(e.get("rule_id"), str) or e.get("severity") not in {"info", "warning", "error"} for e in expected):
            raise ValueError("expected_findings 必须声明 rule_id 和有效 severity")
        if not isinstance(absent, list) or any(not isinstance(r, str) for r in absent):
            raise ValueError("expected_absent_rules 必须是字符串数组")
        if {e["rule_id"] for e in expected} & set(absent):
            raise ValueError("同一规则不能同时标注为应命中和不应命中")
    return value


def _check_case(case: Mapping[str, Any], root: Path) -> dict[str, Any]:
    case_id = str(case.get("id", ""))
    builder_name = str(case.get("builder", ""))
    builder = BUILDERS.get(builder_name)
    if not case_id or builder is None:
        return {
            "id": case_id or builder_name or "<unknown>",
            "passed": False,
            "errors": ["缺少有效 id 或未知 builder"],
            "findings": [],
            "unsupported": [],
            "tp": [],
            "fp": [],
            "fn": [],
            "tn": [],
        }

    input_path = root / f"{case_id}.docx"
    document = Document()
    builder(document)
    document.save(input_path)
    result = audit_document(
        input_path,
        mode=str(case.get("mode", "full")),
        required_sections=case.get("required_sections"),
    )

    errors: list[str] = []
    if result.get("errors"):
        errors.append(f"审计返回 errors={result['errors']!r}")

    findings = result.get("findings", [])
    actual_rule_ids = {item.get("rule_id") for item in findings if item.get("rule_id")}
    actual_defects = {
        item.get("rule_id")
        for item in findings
        if item.get("severity") in DEFECT_SEVERITIES and item.get("rule_id")
    }

    # 1. 预期命中检查
    for expected in case.get("expected_findings", []):
        if not isinstance(expected, Mapping):
            errors.append(f"无效 expected_findings 项：{expected!r}")
            continue
        rule_id = expected.get("rule_id")
        matches = [item for item in findings if item.get("rule_id") == rule_id]
        if not matches:
            errors.append(f"未命中 finding 规则：{rule_id}")
            continue
        expected_severity = expected.get("severity")
        if expected_severity and not any(item.get("severity") == expected_severity for item in matches):
            errors.append(
                f"finding 规则 {rule_id} 未出现期望级别 {expected_severity}"
            )

    # 2. 明确不应命中规则检查
    expected_absent_rules = set(case.get("expected_absent_rules", []))
    for absent_rule in expected_absent_rules:
        if absent_rule in actual_rule_ids:
            errors.append(f"命中了明确不应出现的规则：{absent_rule}")

    # 3. 混淆矩阵 (针对缺陷告警: warning / error)
    expected_defects = {
        item.get("rule_id")
        for item in case.get("expected_findings", [])
        if isinstance(item, Mapping)
        and item.get("severity") in DEFECT_SEVERITIES
        and item.get("rule_id")
    }

    tp_rules = sorted(expected_defects & actual_defects)
    fn_rules = sorted(expected_defects - actual_defects)
    fp_rules = sorted((actual_defects - expected_defects) | (expected_absent_rules & actual_defects))
    tn_rules = sorted(expected_absent_rules - actual_defects)

    if fp_rules:
        errors.append(f"存在非预期的缺陷规则告警 (FP)：{fp_rules}")
    if fn_rules:
        errors.append(f"缺少预期的缺陷规则告警 (FN)：{fn_rules}")

    # 4. 未审计对象检查
    unsupported = result.get("unsupported_objects", [])
    unsupported_types = {item.get("object_type") for item in unsupported}
    for expected_type in case.get("expected_unsupported", []):
        if expected_type not in unsupported_types:
            errors.append(f"未命中 unsupported object 类型：{expected_type}")
    if "expected_unsupported" in case and unsupported_types != set(case["expected_unsupported"]):
        errors.append("未审计对象类型集合与预期不一致")

    # 5. Summary 检查
    for severity, expected_count in dict(case.get("expected_summary", {})).items():
        actual_count = result.get("summary", {}).get(severity)
        if actual_count != expected_count:
            errors.append(f"summary.{severity} 期望 {expected_count}，实际 {actual_count}")

    return {
        "id": case_id,
        "description": str(case.get("description", "")),
        "passed": not errors,
        "summary": result.get("summary", {}),
        "findings": sorted(actual_rule_ids),
        "unsupported": sorted(unsupported_types),
        "tp": tp_rules,
        "fp": fp_rules,
        "fn": fn_rules,
        "tn": tn_rules,
        "errors": errors,
    }


def evaluate(manifest_input: Path | Mapping[str, Any]) -> dict[str, Any]:
    manifest = _load_manifest(manifest_input)
    manifest_name = str(manifest_input) if isinstance(manifest_input, Path) else "<inline-mapping>"
    with tempfile.TemporaryDirectory(prefix="ai-office-controlled-") as temp:
        root = Path(temp)
        cases = [_check_case(case, root) for case in manifest["cases"] if isinstance(case, Mapping)]

    passed = all(bool(item.get("passed")) for item in cases) and len(cases) == len(manifest["cases"])

    total_tp = sum(len(item.get("tp", [])) for item in cases)
    total_fp = sum(len(item.get("fp", [])) for item in cases)
    total_fn = sum(len(item.get("fn", [])) for item in cases)
    total_tn = sum(len(item.get("tn", [])) for item in cases)

    precision = round(total_tp / (total_tp + total_fp), 4) if (total_tp + total_fp) > 0 else None
    recall = round(total_tp / (total_tp + total_fn), 4) if (total_tp + total_fn) > 0 else None

    return {
        "schema_version": "1.1",
        "manifest": manifest_name,
        "case_count": len(cases),
        "passed_count": sum(bool(item.get("passed")) for item in cases),
        "failed_count": sum(not bool(item.get("passed")) for item in cases),
        "metrics": {
            "scope": "warning_and_error_defects",
            "tp": total_tp,
            "fp": total_fp,
            "fn": total_fn,
            "tn": total_tn,
            "precision": precision,
            "recall": recall,
        },
        "disclaimer": "受控合成样例指标仅反映预设测试集对已知确定性规则的覆盖能力，不代表真实文档整体准确率",
        "raw_documents_retained": False,
        "network_used": False,
        "passed": passed,
        "cases": cases,
    }


def render_report(report: Mapping[str, Any], format_name: str) -> str:
    if format_name == "json":
        return json.dumps(report, ensure_ascii=False, indent=2) + "\n"

    metrics = report.get("metrics", {})
    prec_val = metrics.get("precision")
    rec_val = metrics.get("recall")
    prec_str = f"{prec_val * 100:.1f}%" if prec_val is not None else "不适用 (N/A)"
    rec_str = f"{rec_val * 100:.1f}%" if rec_val is not None else "不适用 (N/A)"

    lines = [
        "# 受控 Word 缺陷样例评测报告",
        "",
        f"- 结果：{'通过' if report.get('passed') else '失败'}",
        f"- 样例：{report.get('case_count', 0)}，通过：{report.get('passed_count', 0)}，失败：{report.get('failed_count', 0)}",
        f"- 缺陷规则指标（warning/error，不混入提示）：TP={metrics.get('tp', 0)}, FP={metrics.get('fp', 0)}, FN={metrics.get('fn', 0)}, TN={metrics.get('tn', 0)}",
        f"- 精确率 (Precision)：{prec_str}，召回率 (Recall)：{rec_str}",
        f"- 声明：{report.get('disclaimer', '')}",
        "- 网络：评测脚本不发起网络请求（不提供网络沙箱）；原始变体：临时生成且不保留",
        "",
        "| 样例 | 结果 | 命中规则 | 未审计对象 | 缺陷指标 | 错误 / 偏差 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for case in report.get("cases", []):
        tp_cnt = len(case.get("tp", []))
        fp_cnt = len(case.get("fp", []))
        fn_cnt = len(case.get("fn", []))
        metric_summary = f"TP={tp_cnt}"
        if fp_cnt > 0:
            metric_summary += f", FP={fp_cnt}"
        if fn_cnt > 0:
            metric_summary += f", FN={fn_cnt}"

        lines.append(
            "| {id} | {status} | {findings} | {unsupported} | {metrics} | {errors} |".format(
                id=case.get("id", ""),
                status="通过" if case.get("passed") else "失败",
                findings=", ".join(case.get("findings", [])) or "无",
                unsupported=", ".join(case.get("unsupported", [])) or "无",
                metrics=metric_summary,
                errors="；".join(case.get("errors", [])) or "无",
            )
        )
    return "\n".join(lines) + "\n"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate temporary controlled .docx cases and verify labels."
    )
    parser.add_argument(
        "--manifest",
        default="fixtures/controlled-cases.json",
        help="controlled case manifest",
    )
    parser.add_argument("--format", choices=("json", "markdown", "terminal"), default="terminal")
    parser.add_argument("--output", help="optional report path; never stores generated .docx cases")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    manifest_path = Path(args.manifest).expanduser().resolve(strict=False)
    try:
        report = evaluate(manifest_path)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        report = {
            "schema_version": "1.1",
            "passed": False,
            "case_count": 0,
            "passed_count": 0,
            "failed_count": 1,
            "errors": [str(exc)],
        }
        rendered = render_report(report, "json" if args.format == "json" else "markdown")
        print(rendered, end="")
        return 2

    rendered = render_report(report, "markdown" if args.format in {"markdown", "terminal"} else "json")
    if args.output:
        output_path = Path(args.output).expanduser().resolve(strict=False)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(rendered, encoding="utf-8")
    else:
        print(rendered, end="")
    return 0 if report.get("passed") else 1


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = ["BUILDERS", "DEFECT_SEVERITIES", "build_parser", "evaluate", "main", "render_report"]
