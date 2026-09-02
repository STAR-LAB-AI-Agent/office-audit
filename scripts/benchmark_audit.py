"""Benchmark the deterministic audit path without retaining document files."""

from __future__ import annotations

import argparse
import json
import statistics
import tempfile
import time
from pathlib import Path
from typing import Any, Sequence

from docx import Document

try:
    from .audit_docx import audit_document
except ImportError:  # direct ``python scripts/benchmark_audit.py`` execution
    from audit_docx import audit_document


def _create_fixture(path: Path, paragraph_count: int) -> None:
    document = Document()
    document.add_heading("Benchmark document", level=1)
    for index in range(paragraph_count):
        suffix = " TODO" if index == paragraph_count // 2 else ""
        document.add_paragraph(f"Synthetic paragraph {index}: stable audit input.{suffix}")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "status"
    table.cell(1, 1).text = "complete"
    document.save(path)


def _benchmark_size(paragraph_count: int, repeats: int) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="audit-benchmark-") as directory:
        source = Path(directory) / "synthetic.docx"
        _create_fixture(source, paragraph_count)
        durations: list[float] = []
        result: dict[str, Any] | None = None
        for _ in range(repeats):
            started = time.perf_counter()
            result = audit_document(source, mode="full")
            durations.append((time.perf_counter() - started) * 1000)
        assert result is not None
        retained_chars = sum(
            len(str(item.get("evidence", "")))
            + len(str(item.get("message", "")))
            + len(str(item.get("suggestion", "")))
            for item in result["findings"]
        )
        input_chars = int(result["metrics"].get("input_chars_seen", 0))
        return {
            "paragraphs": paragraph_count,
            "repeats": repeats,
            "p50_elapsed_ms": round(statistics.median(durations), 2),
            "p95_approx_elapsed_ms": round(max(durations), 2),
            "input_chars_seen": input_chars,
            "deterministic_summary_chars": retained_chars,
            "summary_char_ratio": round(retained_chars / input_chars, 4) if input_chars else 0.0,
            "model_calls": result["metrics"].get("model_calls", 0),
            "findings": len(result["findings"]),
            "errors": len(result["errors"]),
        }


def _sizes(value: str) -> list[int]:
    parsed = [int(item.strip()) for item in value.split(",") if item.strip()]
    if not parsed or any(item <= 0 for item in parsed):
        raise argparse.ArgumentTypeError("sizes must be positive comma-separated integers")
    return parsed


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Benchmark the local deterministic .docx audit path using temporary fixtures."
    )
    parser.add_argument("--sizes", type=_sizes, default=[50, 200, 800])
    parser.add_argument("--repeats", type=int, default=3)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.repeats <= 0:
        raise SystemExit("--repeats must be positive")
    report = {
        "benchmark": "deterministic_docx_audit",
        "note": "summary_char_ratio is a local evidence-size proxy, not a model token count",
        "results": [_benchmark_size(size, args.repeats) for size in args.sizes],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = ["build_parser", "main"]
