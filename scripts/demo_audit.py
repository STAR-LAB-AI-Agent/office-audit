"""Create a safe synthetic demo document and run the local audit once."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

try:
    from .audit_docx import audit_document, classify_request, render_result
except ImportError:  # direct ``python scripts/demo_audit.py`` execution
    from audit_docx import audit_document, classify_request, render_result


def _create_demo_document(path: Path) -> None:
    document = Document()
    document.add_heading("文档审计演示", level=1)
    document.add_heading("直接跳到第三级标题", level=3)
    document.add_paragraph("TODO：请在提交前补充本字段。")
    empty_paragraph = document.add_paragraph()
    empty_paragraph.add_run("格式示例")
    empty_paragraph.runs[0].font.size = Pt(12)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "备注"
    # Keep the bottom-right cell empty to exercise the empty-cell rule.
    drawing_paragraph = document.add_paragraph()
    drawing_paragraph._p.get_or_add_pPr()
    drawing_paragraph._p.append(OxmlElement("w:drawing"))
    document.save(path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Create a synthetic .docx and run a safe local audit for demonstrations."
    )
    parser.add_argument(
        "--output-dir",
        default="outputs/demo",
        help="directory for the generated demo input and reports",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="replace existing demo files in the selected output directory",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    output_dir = Path(args.output_dir).expanduser().resolve(strict=False)
    output_dir.mkdir(parents=True, exist_ok=True)
    input_path = output_dir / "demo-input.docx"
    json_path = output_dir / "demo-result.json"
    markdown_path = output_dir / "demo-result.md"
    targets = (input_path, json_path, markdown_path)
    if not args.force and any(path.exists() for path in targets):
        print(json.dumps({"error": "demo_output_exists", "output_dir": str(output_dir)}, ensure_ascii=False))
        return 2

    _create_demo_document(input_path)
    request = "全面检查这份 Word 文档"
    routing = classify_request(request)
    if routing.get("status") != "ok":
        print(json.dumps({"error": "demo_request_routing_failed", "routing": routing}, ensure_ascii=False))
        return 2
    result = audit_document(input_path, mode=routing["mode"])
    result["request_routing"] = {
        "request": request,
        "intent": routing.get("intent"),
        "selected_mode": routing.get("mode"),
        "matched_terms": routing.get("matched_terms", []),
    }
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    markdown_path.write_text(render_result(result, "markdown"), encoding="utf-8")
    summary = result["summary"]
    print(
        json.dumps(
            {
                "input": str(input_path),
                "json_report": str(json_path),
                "markdown_report": str(markdown_path),
                "mode": result["mode"],
                "summary": summary,
                "unsupported_objects": len(result["unsupported_objects"]),
                "model_calls": result["metrics"]["model_calls"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = ["build_parser", "main"]
