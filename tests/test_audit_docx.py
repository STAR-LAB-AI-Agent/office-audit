from __future__ import annotations

import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from scripts.audit_docx import audit_document, classify_request, main, render_result


class AuditDocxTests(unittest.TestCase):
    def test_markdown_findings_include_human_readable_paragraph_context(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "location-context.docx"
            document = Document()
            document.add_paragraph("上文定位文字")
            document.add_paragraph("")
            document.add_paragraph("")
            document.add_paragraph("下文定位文字")
            document.add_paragraph("同类正文一")
            document.add_paragraph("同类正文二")
            outlier = document.add_paragraph("格式离群段落")
            outlier.runs[0].bold = True
            document.save(path)

            result = audit_document(path, mode="fields_format")
            markdown = render_result(result, "markdown")

            self.assertIn(
                "位置：正文中“上文定位文字”之后、“下文定位文字”之前的第1个空白段落（连续2个）",
                markdown,
            )
            self.assertIn(
                "位置：正文中“上文定位文字”之后、“下文定位文字”之前的第2个空白段落（连续2个）",
                markdown,
            )
            self.assertIn("前文：“上文定位文字”", markdown)
            self.assertIn("后文：“下文定位文字”", markdown)
            self.assertIn("位置：正文中以“格式离群段落”开头的段落", markdown)
            self.assertIn("段落文字：“格式离群段落”", markdown)
            self.assertIn("格式差异：", markdown)
            self.assertNotIn("- 位置：正文第", markdown)

    def test_structural_header_footer_blanks_are_not_reported_as_content_blanks(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "structural-header-footer.docx"
            document = Document()
            document.add_paragraph("正文")
            section = document.sections[0]
            section.different_first_page_header_footer = True
            for story in (
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ):
                _ = story.paragraphs[0]
            document.save(path)

            result = audit_document(path, mode="fields_format")
            structural_blanks = [
                finding
                for finding in result["findings"]
                if finding["rule_id"] == "fields.empty_paragraph"
                and finding["location"].get("part") in {"header", "footer"}
            ]

            self.assertEqual([], structural_blanks)

    def _write_sample(self, directory: Path, *, with_drawing: bool = False) -> Path:
        path = directory / "sample.docx"
        document = Document()
        document.add_heading("项目说明", level=1)
        document.add_heading("实现细节", level=3)
        document.add_paragraph("这是一段用于回归测试的正文。")
        document.add_paragraph("TODO：补充审计范围")
        document.add_paragraph("")
        for text in ("统一段落一", "统一段落二", "统一段落三", "格式离群段落"):
            paragraph = document.add_paragraph(text)
            if text == "格式离群段落":
                paragraph.runs[0].bold = True
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = ""
        table.cell(1, 0).text = "值"
        table.cell(1, 1).text = "待填写"
        document.sections[0].header.paragraphs[0].text = "测试页眉"
        document.sections[0].footer.paragraphs[0].text = "测试页脚"
        if with_drawing:
            paragraph = document.add_paragraph()
            run = paragraph.add_run()
            run._r.append(OxmlElement("w:drawing"))
        document.save(path)
        return path

    def test_full_mode_returns_schema_and_stable_findings(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp))
            before = source.read_bytes()
            result = audit_document(source, mode="full")

            self.assertEqual(result["schema_version"], "0.1")
            self.assertEqual(result["audit_target"]["format"], "docx")
            self.assertTrue(result["scope"]["unsupported_detection"])
            self.assertEqual(result["errors"], [])
            self.assertEqual(
                result["summary"]["warning"],
                sum(item["severity"] == "warning" for item in result["findings"]),
            )
            rule_ids = {item["rule_id"] for item in result["findings"]}
            self.assertIn("structure.heading_level_jump", rule_ids)
            self.assertIn("fields.placeholder", rule_ids)
            self.assertIn("fields.empty_cell", rule_ids)
            self.assertIn("format.dominant_style_outlier", rule_ids)
            self.assertEqual(
                [item["id"] for item in result["findings"]],
                [f"F-{index:03d}" for index in range(1, len(result["findings"]) + 1)],
            )
            self.assertFalse(
                any(
                    item["location"].get("variant") in {"first_page", "even_page"}
                    for item in result["findings"]
                )
            )
            self.assertEqual(before, source.read_bytes())

    def test_structure_mode_checks_user_sections_without_guessing(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp))
            result = audit_document(
                source,
                mode="structure",
                required_sections=["项目说明", "风险与建议"],
            )

            self.assertTrue(
                any(
                    item["rule_id"] == "structure.required_section_missing"
                    and item["severity"] == "error"
                    for item in result["findings"]
                )
            )
            self.assertFalse(any(item["rule_id"].startswith("fields.") for item in result["findings"]))
            self.assertFalse(any(item["rule_id"].startswith("format.") for item in result["findings"]))
            self.assertNotIn("structure.required_sections_not_specified", {item["rule_id"] for item in result["findings"]})

    def test_fields_format_mode_excludes_structure_rules(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp))
            result = audit_document(source, mode="fields_format")
            rule_ids = {item["rule_id"] for item in result["findings"]}

            self.assertIn("fields.placeholder", rule_ids)
            self.assertIn("format.dominant_style_outlier", rule_ids)
            self.assertFalse(any(rule_id.startswith("structure.") for rule_id in rule_ids))

    def test_custom_placeholder_and_severity_rules_are_applied(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp))
            document = Document(source)
            document.add_paragraph("未完成：需要人工确认")
            document.save(source)
            result = audit_document(
                source,
                mode="fields_format",
                rules={
                    "placeholders": ["未完成"],
                    "severity_overrides": {"fields.placeholder": "error"},
                },
            )

            placeholder_findings = [
                item for item in result["findings"] if item["rule_id"] == "fields.placeholder"
            ]
            self.assertEqual(len(placeholder_findings), 1)
            self.assertEqual(placeholder_findings[0]["severity"], "error")
            self.assertNotIn("TODO", placeholder_findings[0]["evidence"])

    def test_unsupported_objects_are_individually_described(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp), with_drawing=True)
            result = audit_document(source)
            unsupported = result["unsupported_objects"]

            self.assertTrue(unsupported)
            drawing = next(item for item in unsupported if item["object_type"] == "image_or_drawing")
            self.assertGreaterEqual(drawing["count"], 1)
            self.assertTrue(drawing["location"].get("xml_path"))
            for field in ("reason", "impact", "suggestion"):
                self.assertTrue(drawing[field])
            markdown = render_result(result, "markdown")
            self.assertIn("## 未审计对象", markdown)
            self.assertIn(drawing["id"], markdown)

    def test_invalid_input_is_structured_without_traceback(self) -> None:
        missing = Path(tempfile.gettempdir()) / "audit-input-that-does-not-exist.docx"
        result = audit_document(missing)
        rendered = render_result(result, "json")

        self.assertEqual(result["errors"][0]["code"], "input_not_found")
        self.assertEqual(result["summary"]["error"], 1)
        self.assertNotIn("Traceback", rendered)

    def test_cli_writes_a_separate_report_and_keeps_source_unchanged(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            directory = Path(temp)
            source = self._write_sample(directory)
            report = directory / "reports" / "result.md"
            log = directory / "logs" / "audit.jsonl"
            before = source.read_bytes()
            exit_code = main(
                [
                    "--input",
                    str(source),
                    "--request",
                    "结构审计，备注 sk-test-secret",
                    "--mode",
                    "structure",
                    "--format",
                    "markdown",
                    "--output",
                    str(report),
                    "--log",
                    str(log),
                ]
            )

            self.assertEqual(exit_code, 0)
            self.assertTrue(report.is_file())
            self.assertIn("# AI Office 文档审计", report.read_text(encoding="utf-8"))
            log_text = log.read_text(encoding="utf-8")
            self.assertIn('"event":"audit_completed"', log_text)
            self.assertNotIn("TODO", log_text)
            self.assertNotIn("evidence", log_text)
            self.assertNotIn("sk-test-secret", log_text)
            self.assertEqual(before, source.read_bytes())

    def test_cli_rejects_report_path_that_overwrites_source(self) -> None:
        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp))
            before = source.read_bytes()
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                exit_code = main(["--input", str(source), "--output", str(source)])

            result = json.loads(output.getvalue())
            self.assertEqual(exit_code, 2)
            self.assertEqual(result["errors"][0]["code"], "unsafe_output_path")
            self.assertEqual(before, source.read_bytes())

    def test_natural_language_router_maps_three_intents_and_runs_audit(self) -> None:
        cases = (
            ("全面检查这份 Word 文档", "full"),
            ("只看标题层级和缺失章节", "structure"),
            ("检查空字段、占位符和格式是否统一", "fields_format"),
        )
        for request, expected_mode in cases:
            routing = classify_request(request)
            self.assertEqual(routing["status"], "ok")
            self.assertEqual(routing["mode"], expected_mode)

        with tempfile.TemporaryDirectory(prefix="audit-test-") as temp:
            source = self._write_sample(Path(temp))
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                exit_code = main(
                    [
                        "--input",
                        str(source),
                        "--request",
                        "只看标题层级和缺失章节",
                    ]
                )
            result = json.loads(output.getvalue())
            self.assertEqual(exit_code, 0)
            self.assertEqual(result["mode"], "structure")
            self.assertEqual(result["request_routing"]["selected_mode"], "structure")

    def test_natural_language_router_rejects_ambiguous_or_unknown_requests(self) -> None:
        ambiguous = classify_request("只检查标题和格式")
        unknown = classify_request("帮我看看这份材料")
        self.assertEqual(ambiguous["error"]["code"], "ambiguous_intent")
        self.assertEqual(unknown["error"]["code"], "unknown_intent")


if __name__ == "__main__":
    unittest.main()
