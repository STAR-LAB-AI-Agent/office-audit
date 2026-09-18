"""Tests for Phase 3 configuration validation and safety resource limits."""
import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from scripts.audit_docx import (
    AuditInputError,
    audit_document,
    main,
    validate_rules_configuration,
)


class Phase3RulesAndLimitsTests(unittest.TestCase):
    def _create_sample(self, folder: Path) -> Path:
        doc = Document()
        doc.add_heading("第一章 概述", level=1)
        doc.add_paragraph("这是正常的一段正文。")
        doc.add_paragraph("")  # empty paragraph
        for text in ("正文1", "正文2", "正文3", "离群段落"):
            p = doc.add_paragraph(text)
            if text == "离群段落":
                p.runs[0].bold = True
        path = folder / "sample.docx"
        doc.save(path)
        return path

    def test_unknown_rule_config_key_raises_invalid_rules(self):
        with self.assertRaises(AuditInputError) as ctx:
            validate_rules_configuration({"unknown_custom_key": True})
        self.assertEqual(ctx.exception.code, "invalid_rules")

        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            res = audit_document(source, rules={"unknown_custom_key": True})
            self.assertEqual(res["errors"][0]["code"], "invalid_rules")

            # Also verify CLI exit code 2
            rules_path = Path(temp) / "bad_rules.json"
            rules_path.write_text(json.dumps({"unknown_key": 123}), encoding="utf-8")
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                code = main(["--input", str(source), "--rules", str(rules_path)])
            self.assertEqual(code, 2)
            cli_res = json.loads(output.getvalue())
            self.assertEqual(cli_res["errors"][0]["code"], "invalid_rules")

    def test_unknown_rule_id_raises_invalid_rules(self):
        with self.assertRaises(AuditInputError) as ctx:
            validate_rules_configuration({"severity_overrides": {"nonexistent.rule": "error"}})
        self.assertEqual(ctx.exception.code, "invalid_rules")

        with self.assertRaises(AuditInputError) as ctx:
            validate_rules_configuration({"disabled_rules": ["nonexistent.rule"]})
        self.assertEqual(ctx.exception.code, "invalid_rules")

        with self.assertRaises(AuditInputError) as ctx:
            validate_rules_configuration({"rule_switches": {"nonexistent.rule": False}})
        self.assertEqual(ctx.exception.code, "invalid_rules")

        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            res = audit_document(source, rules={"severity_overrides": {"nonexistent.rule": "error"}})
            self.assertEqual(res["errors"][0]["code"], "invalid_rules")

    def test_invalid_severity_raises_invalid_rules(self):
        with self.assertRaises(AuditInputError) as ctx:
            validate_rules_configuration({"severity_overrides": {"fields.empty_paragraph": "fatal"}})
        self.assertEqual(ctx.exception.code, "invalid_rules")

        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            res = audit_document(source, rules={"severity_overrides": {"fields.empty_paragraph": "fatal"}})
            self.assertEqual(res["errors"][0]["code"], "invalid_rules")

    def test_invalid_placeholder_regex_raises_invalid_rules(self):
        with self.assertRaises(AuditInputError) as ctx:
            validate_rules_configuration({"placeholder_pattern": "(unclosed_group"})
        self.assertEqual(ctx.exception.code, "invalid_rules")

        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            res = audit_document(source, rules={"placeholder_pattern": "(unclosed_group"})
            self.assertEqual(res["errors"][0]["code"], "invalid_rules")

    def test_disabling_rules_via_disabled_rules_and_switches(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))

            # Default run hits empty_paragraph and dominant_style_outlier
            base_result = audit_document(source)
            base_rules = {f["rule_id"] for f in base_result["findings"]}
            self.assertIn("fields.empty_paragraph", base_rules)
            self.assertIn("format.dominant_style_outlier", base_rules)

            # Disable via disabled_rules
            disabled_result = audit_document(
                source,
                rules={"disabled_rules": ["format.dominant_style_outlier", "fields.empty_paragraph"]},
            )
            disabled_rules = {f["rule_id"] for f in disabled_result["findings"]}
            self.assertNotIn("format.dominant_style_outlier", disabled_rules)
            self.assertNotIn("fields.empty_paragraph", disabled_rules)

            # Disable via rule_switches
            switched_result = audit_document(
                source,
                rules={"rule_switches": {"format.dominant_style_outlier": False}},
            )
            switched_rules = {f["rule_id"] for f in switched_result["findings"]}
            self.assertNotIn("format.dominant_style_outlier", switched_rules)
            self.assertIn("fields.empty_paragraph", switched_rules)

    def test_max_file_size_limit_triggers_file_too_large(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            # Set tiny max_file_size_bytes
            res = audit_document(source, rules={"limits": {"max_file_size_bytes": 10}})
            self.assertEqual(res["errors"][0]["code"], "file_too_large")

            # CLI verification
            rules_path = Path(temp) / "limits.json"
            rules_path.write_text(json.dumps({"limits": {"max_file_size_bytes": 10}}), encoding="utf-8")
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                code = main(["--input", str(source), "--rules", str(rules_path)])
            self.assertEqual(code, 2)
            cli_res = json.loads(output.getvalue())
            self.assertEqual(cli_res["errors"][0]["code"], "file_too_large")

    def test_zip_limits_trigger_zip_bomb_detected(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            # Limit zip entries
            res = audit_document(source, rules={"limits": {"max_zip_entries": 1}})
            self.assertEqual(res["errors"][0]["code"], "zip_bomb_detected")

            # Limit uncompressed bytes
            res = audit_document(source, rules={"limits": {"max_uncompressed_bytes": 100}})
            self.assertEqual(res["errors"][0]["code"], "zip_bomb_detected")

            # CLI verification
            rules_path = Path(temp) / "zip_limits.json"
            rules_path.write_text(json.dumps({"limits": {"max_zip_entries": 1}}), encoding="utf-8")
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                code = main(["--input", str(source), "--rules", str(rules_path)])
            self.assertEqual(code, 2)
            cli_res = json.loads(output.getvalue())
            self.assertEqual(cli_res["errors"][0]["code"], "zip_bomb_detected")

    def test_rules_example_json_is_valid(self):
        example_path = Path("references/rules-example.json")
        self.assertTrue(example_path.is_file())
        rules = json.loads(example_path.read_text(encoding="utf-8"))
        validated = validate_rules_configuration(rules)
        self.assertEqual(validated["limits"]["max_file_size_bytes"], 52428800)

        with tempfile.TemporaryDirectory() as temp:
            source = self._create_sample(Path(temp))
            result = audit_document(source, rules=rules)
            self.assertEqual(result["errors"], [])
            # Format outlier is disabled in rules-example.json
            rule_ids = {f["rule_id"] for f in result["findings"]}
            self.assertNotIn("format.dominant_style_outlier", rule_ids)


if __name__ == "__main__":
    unittest.main()
