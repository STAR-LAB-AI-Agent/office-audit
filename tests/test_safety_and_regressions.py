"""Regression tests for defects found in the delivery review."""
import contextlib
import hashlib
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from scripts.audit_docx import AuditInputError, audit_document, main, render_result, validate_rules_configuration
from scripts.evaluate_controlled_cases import evaluate, main as evaluate_main


class DeliveryReviewTests(unittest.TestCase):
    def test_resource_limits_reject_nonfinite_boolean_and_fractional_counts(self):
        for value in (True, False, float('nan'), float('inf'), -1, 0, 1.5):
            with self.subTest(value=value), self.assertRaises(AuditInputError):
                validate_rules_configuration({'limits': {'max_zip_entries': value}})
        validate_rules_configuration({'limits': {'max_compression_ratio': 1.5}})

    def test_manifest_rejects_path_ids_empty_cases_duplicates_and_conflicts(self):
        valid = {'id': 'valid', 'builder': 'clean_general_document'}
        bad_cases = [[], [dict(valid, id='../escape')], [dict(valid, id='C:\\escape')], [valid, valid], [None],
                     [dict(valid, expected_findings=[{'rule_id': 'x', 'severity': 'warning'}], expected_absent_rules=['x'])]]
        for cases in bad_cases:
            with self.subTest(cases=cases), self.assertRaises(ValueError):
                evaluate({'cases': cases})

    def test_manifest_validation_occurs_before_document_creation(self):
        with patch('scripts.evaluate_controlled_cases.Document') as build:
            with self.assertRaises(ValueError):
                evaluate({'cases': [{'id': '../escape', 'builder': 'clean_general_document'}]})
            build.assert_not_called()

    def test_missing_unexpected_objects_fails_evaluation(self):
        report = evaluate({'cases': [{'id': 'objects', 'builder': 'unsupported_drawing_object', 'expected_unsupported': []}]})
        self.assertFalse(report['passed'])

    def test_cli_reports_fp_and_fn_as_exit_one(self):
        for case in (
            {'id': 'fp', 'builder': 'table_cell_outlier', 'mode': 'fields_format', 'expected_findings': []},
            {'id': 'fn', 'builder': 'clean_general_document', 'expected_findings': [{'rule_id': 'fields.placeholder', 'severity': 'warning'}]},
        ):
            with tempfile.TemporaryDirectory() as temp:
                manifest = Path(temp) / 'labels.json'
                manifest.write_text(json.dumps({'cases': [case]}), encoding='utf-8')
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(evaluate_main(['--manifest', str(manifest)]), 1)

    def test_equations_and_image_only_cells_are_located_not_empty(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'objects.docx'
            doc = Document()
            doc.add_heading('对象', level=1)
            math_para = OxmlElement('m:oMathPara')
            math_para.append(OxmlElement('m:oMath'))
            doc.add_paragraph()._p.append(math_para)
            doc.add_paragraph()._p.append(OxmlElement('m:oMath'))
            cell = doc.add_table(rows=1, cols=1).cell(0, 0)
            cell.paragraphs[0].add_run()._r.append(OxmlElement('w:drawing'))
            doc.save(source)
            result = audit_document(source)
            self.assertFalse(result['errors'])
            self.assertNotIn('fields.empty_cell', {f['rule_id'] for f in result['findings']})
            equations = [o for o in result['unsupported_objects'] if o['object_type'] == 'equation']
            self.assertEqual(len(equations), 2)
            self.assertTrue(all('paragraph_index' in e['location'] and e.get('locator') for e in equations))

    def test_disabled_rules_are_visible_and_skip_placeholder_search(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'plain.docx'
            doc = Document()
            doc.add_paragraph('TODO')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'TODO'
            doc.save(source)
            disabled = ['fields.placeholder', 'fields.placeholder_cell']
            with patch('scripts.audit_docx._placeholder_regex') as regex:
                result = audit_document(source, rules={'disabled_rules': disabled})
                regex.return_value.search.assert_not_called()
            self.assertEqual(result['disabled_rules'], disabled)
            self.assertIn('本次未执行的规则', render_result(result, 'markdown'))

    def test_hardlink_output_cannot_overwrite_source(self):
        with tempfile.TemporaryDirectory() as temp:
            source, alias = Path(temp) / 'source.docx', Path(temp) / 'alias.docx'
            Document().save(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            try:
                os.link(source, alias)
            except OSError as exc:
                self.skipTest(str(exc))
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main(['--input', str(source), '--output', str(alias)]), 2)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
