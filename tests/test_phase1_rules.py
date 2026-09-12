"""Regression cases for intentional layout and real formatting defects."""
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from scripts.audit_docx import audit_document, render_result


class Phase1RulesTests(unittest.TestCase):
    def audit(self, document, **kwargs):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'sample.docx'
            document.save(path)
            before = path.read_bytes()
            result = audit_document(path, **kwargs)
            self.assertEqual(before, path.read_bytes())
            self.assertEqual(result['errors'], [])
            return result

    def findings(self, result, rule):
        return [f for f in result['findings'] if f['rule_id'] == rule]

    def test_single_blank_is_info_and_override_still_works(self):
        doc = Document()
        doc.add_paragraph('上文')
        doc.add_paragraph()
        doc.add_paragraph('下文')
        result = self.audit(doc)
        self.assertEqual(self.findings(result, 'fields.empty_paragraph')[0]['severity'], 'info')
        result = self.audit(doc, rules={'severity_overrides': {'fields.empty_paragraph': 'error'}})
        self.assertEqual(self.findings(result, 'fields.empty_paragraph')[0]['severity'], 'error')

    def test_blank_groups_do_not_cross_tables(self):
        doc = Document()
        doc.add_paragraph()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格'
        doc.add_paragraph()
        blanks = self.findings(self.audit(doc), 'fields.empty_paragraph')
        self.assertEqual([f['location']['paragraph_indices'] for f in blanks], [[0], [1]])
        self.assertTrue(all(f['severity'] == 'info' for f in blanks))

    def test_content_carriers_are_not_blank(self):
        doc = Document()
        doc.add_paragraph('对象上文')
        doc.add_paragraph().add_run()._r.append(OxmlElement('w:drawing'))
        doc.add_paragraph()._p.append(OxmlElement('m:oMathPara'))
        doc.add_page_break()
        doc.add_paragraph('对象下文')
        result = self.audit(doc)
        self.assertEqual(self.findings(result, 'fields.empty_paragraph'), [])
        self.assertEqual(len(result['unsupported_objects']), 1)
        self.assertIn('对象上文', render_result(result, 'markdown'))

    def test_title_candidate_is_not_a_verified_required_section(self):
        doc = Document()
        doc.add_paragraph('项目说明', 'Title')
        result = self.audit(doc, required_sections=['项目说明'])
        self.assertEqual(len(self.findings(result, 'structure.title_candidate')), 1)
        self.assertEqual(len(self.findings(result, 'structure.required_section_missing')), 1)
        self.assertIn('不表示页面上没有标题', self.findings(result, 'structure.no_headings')[0]['evidence'])

    def test_plain_opening_is_not_a_title_candidate(self):
        doc = Document()
        doc.add_paragraph('普通正文')
        self.assertEqual(self.findings(self.audit(doc), 'structure.title_candidate'), [])

    def test_roles_do_not_hide_a_real_body_outlier(self):
        doc = Document()
        for i in range(4):
            doc.add_paragraph('普通正文' + str(i))
        doc.add_paragraph('正文异常').runs[0].bold = True
        doc.add_paragraph('图1 示例').runs[0].italic = True
        doc.add_paragraph('引用文字', 'Quote')
        doc.add_paragraph('列表文字', 'List Bullet')
        doc.add_paragraph('[1] Reference entry').runs[0].italic = True
        outliers = self.findings(self.audit(doc), 'format.dominant_style_outlier')
        self.assertEqual(len(outliers), 1)
        self.assertEqual(outliers[0]['locator']['text'], '正文异常')

    def test_no_majority_means_no_dominant_style_warning(self):
        doc = Document()
        for i in range(4):
            p = doc.add_paragraph('正文' + str(i))
            if i > 1:
                p.runs[0].bold = True
        self.assertEqual(self.findings(self.audit(doc), 'format.dominant_style_outlier'), [])

    def test_table_columns_have_separate_baselines(self):
        doc = Document()
        table = doc.add_table(rows=5, cols=2)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = '内容'
                if j == 1 or i == 0:
                    cell.paragraphs[0].runs[0].bold = True
        self.assertEqual(self.findings(self.audit(doc), 'format.dominant_style_outlier'), [])
        table.cell(3, 0).paragraphs[0].runs[0].italic = True
        self.assertEqual(len(self.findings(self.audit(doc), 'format.dominant_style_outlier')), 1)

    def test_objects_in_cells_and_headers_keep_context(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = '对象所在单元格'
        cell.paragraphs[0].add_run()._r.append(OxmlElement('w:drawing'))
        header = doc.sections[0].header.paragraphs[0]
        header.text = '页眉对象'
        header.add_run()._r.append(OxmlElement('w:drawing'))
        result = self.audit(doc)
        objects = result['unsupported_objects']
        self.assertEqual(len(objects), 2)
        self.assertEqual(objects[0]['location']['cell_index'], 0)
        self.assertEqual(objects[1]['location']['part'], 'header')
        self.assertTrue(all(o['location']['xml_path'] and o['locator']['text'] for o in objects))


if __name__ == '__main__':
    unittest.main()
